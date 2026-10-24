import os
from typing import List, Tuple
from app.core.errors import AppError
from app.core.logging import logger

class ExtractedPage:
    def __init__(self, page_number: int, text: str):
        self.page_number = page_number
        self.text = text

class DocumentExtractionResult:
    def __init__(
        self,
        pages: List[ExtractedPage],
        file_type: str,
        is_scanned: bool = False,
        warning: str | None = None
    ):
        self.pages = pages
        self.file_type = file_type
        self.is_scanned = is_scanned
        self.warning = warning

    @property
    def total_char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)

    @property
    def page_count(self) -> int:
        return len(self.pages)

def extract_text_from_file(file_path: str, max_size_mb: int = 20) -> DocumentExtractionResult:
    """
    Extracts text from PDF, DOCX, or TXT file preserving page metadata where possible.
    Includes validation, encoding fallback, and scanned PDF detection.
    """
    if not os.path.exists(file_path):
        raise AppError("FILE_NOT_FOUND", f"File not found: {file_path}", status_code=404)

    file_size_bytes = os.path.getsize(file_path)
    if file_size_bytes > max_size_mb * 1024 * 1024:
        raise AppError(
            "FILE_TOO_LARGE",
            f"File exceeds maximum allowed size of {max_size_mb} MB ({file_size_bytes} bytes)",
            status_code=413
        )

    _, ext = os.path.splitext(file_path)
    ext = ext.lower().lstrip(".")

    if ext == "pdf":
        return _extract_pdf(file_path)
    elif ext == "docx":
        return _extract_docx(file_path)
    elif ext in ["txt", "md", "csv"]:
        return _extract_txt(file_path)
    else:
        raise AppError(
            "UNSUPPORTED_FORMAT",
            f"Unsupported file format '{ext}'. Supported formats: PDF, DOCX, TXT.",
            status_code=415
        )

def _extract_pdf(file_path: str) -> DocumentExtractionResult:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise AppError("DEPENDENCY_MISSING", "PyMuPDF is not installed.", status_code=500)

    try:
        doc = fitz.open(file_path)
    except Exception as e:
        raise AppError("CORRUPT_FILE", f"Unable to parse PDF file: {str(e)}", status_code=400)

    pages: List[ExtractedPage] = []
    total_chars = 0

    try:
        for idx, page in enumerate(doc):
            text = page.get_text("text") or ""
            cleaned = text.strip()
            pages.append(ExtractedPage(page_number=idx + 1, text=cleaned))
            total_chars += len(cleaned)
    finally:
        doc.close()

    if not pages:
        pages.append(ExtractedPage(page_number=1, text=""))

    # Scanned PDF detection
    is_scanned = False
    warning = None
    if total_chars < 30 and len(pages) > 0:
        is_scanned = True
        warning = (
            "Notice: Minimal or no text extracted. This document appears to be a scanned or image-only PDF. "
            "OCR (Optical Character Recognition) is NOT supported."
        )
        logger.warning(f"Scanned PDF detected for {file_path}")

    return DocumentExtractionResult(pages=pages, file_type="pdf", is_scanned=is_scanned, warning=warning)

def _extract_docx(file_path: str) -> DocumentExtractionResult:
    try:
        import docx
    except ImportError:
        raise AppError("DEPENDENCY_MISSING", "python-docx is not installed.", status_code=500)

    try:
        doc = docx.Document(file_path)
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text)
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))

        combined = "\n\n".join(full_text)
        # DOCX format does not have explicit physical page breaks without rendering; single page baseline
        pages = [ExtractedPage(page_number=1, text=combined)]
        return DocumentExtractionResult(pages=pages, file_type="docx")
    except Exception as e:
        raise AppError("CORRUPT_FILE", f"Unable to parse DOCX file: {str(e)}", status_code=400)

def _extract_txt(file_path: str) -> DocumentExtractionResult:
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]
    content = None

    with open(file_path, "rb") as f:
        raw_bytes = f.read()

    for enc in encodings:
        try:
            content = raw_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue

    if content is None:
        content = raw_bytes.decode("utf-8", errors="replace")

    # If the text file has custom page headers like "--- PAGE 1 ---", preserve pages!
    lines = content.splitlines()
    pages: List[ExtractedPage] = []
    current_page = 1
    current_lines: List[str] = []

    for line in lines:
        if line.strip().startswith("--- PAGE ") and line.strip().endswith(" ---"):
            if current_lines:
                pages.append(ExtractedPage(page_number=current_page, text="\n".join(current_lines).strip()))
                current_lines = []
            try:
                page_str = line.strip().replace("--- PAGE ", "").replace(" ---", "")
                current_page = int(page_str)
            except ValueError:
                current_page += 1
        else:
            current_lines.append(line)

    if current_lines or not pages:
        pages.append(ExtractedPage(page_number=current_page, text="\n".join(current_lines).strip()))

    return DocumentExtractionResult(pages=pages, file_type="txt")
